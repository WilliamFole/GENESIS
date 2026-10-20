import csv
import io
from calendar import monthrange
from datetime import date
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.models.cliente import Cliente
from app.models.documento import Documento, EstadoDocumento, TipoDocumentoFiscal
from app.models.pagamento import Pagamento
from app.models.serie import Serie

MESES_PT = [
    "janeiro", "fevereiro", "março", "abril", "maio", "junho",
    "julho", "agosto", "setembro", "outubro", "novembro", "dezembro",
]


def _ultimo_dia(d: date) -> int:
    return monthrange(d.year, d.month)[1]


def nome_extrato(tipo: str, data_inicio: date, data_fim: date, ext: str, prefixo: str = "Extrato de") -> str:
    """Gera nome de ficheiro legivel com base no periodo."""
    ext = f".{ext}" if ext else ""
    if (
        data_inicio.year == data_fim.year
        and data_inicio.month == data_fim.month
        and data_inicio.day == 1
        and data_fim.day == _ultimo_dia(data_fim)
    ):
        return (
            f"{prefixo} {tipo} do mes de "
            f"{MESES_PT[data_inicio.month - 1]} {data_inicio.year}{ext}"
        )
    if (
        data_inicio.year == data_fim.year
        and data_inicio.day == 1
        and data_fim.day == _ultimo_dia(data_fim)
    ):
        return (
            f"{prefixo} {tipo} de "
            f"{MESES_PT[data_inicio.month - 1]} a "
            f"{MESES_PT[data_fim.month - 1]} {data_inicio.year}{ext}"
        )
    return (
        f"{prefixo} {tipo} de "
        f"({data_inicio.strftime('%d-%m-%Y')} a "
        f"{data_fim.strftime('%d-%m-%Y')}){ext}"
    )


def _calcular_credito(db: Session, doc_id: int) -> Decimal:
    valor = db.query(func.coalesce(func.sum(Documento.total), 0)).filter(
        Documento.documento_origem_id == doc_id,
        Documento.tipo_documento == TipoDocumentoFiscal.NC,
        Documento.estado == EstadoDocumento.FINALIZADO,
    ).scalar()
    return Decimal(valor)


def relatorio_vendas(
    db: Session,
    data_inicio: date,
    data_fim: date,
    cliente_id: int | None = None,
    artigo_id: int | None = None,
):
    q = db.query(Documento).filter(
        Documento.data_documento >= data_inicio,
        Documento.data_documento <= data_fim,
        Documento.estado == EstadoDocumento.FINALIZADO,
        Documento.tipo_documento.in_([
            TipoDocumentoFiscal.FT,
            TipoDocumentoFiscal.VD,
            TipoDocumentoFiscal.FS,
            TipoDocumentoFiscal.NC,
        ]),
    )
    if cliente_id:
        q = q.filter(Documento.cliente_id == cliente_id)

    docs = q.all()
    bruto = Decimal("0")
    liquido = Decimal("0")
    total_iva = Decimal("0")
    por_cliente: dict[int, dict] = {}

    for d in docs:
        if artigo_id and not any(l.artigo_id == artigo_id for l in d.linhas):
            continue
        is_nc = d.tipo_documento == TipoDocumentoFiscal.NC
        if not is_nc:
            bruto += d.total
            liquido += d.total
            total_iva += d.total_iva
        else:
            liquido -= d.total
            total_iva -= d.total_iva

            # ... (dentro do for d in docs:)
        cid = d.cliente_id or 0
        if cid not in por_cliente:
            por_cliente[cid] = {"nome": d.cliente_nome or "Consumidor Final", "bruto": Decimal("0"), "liquido": Decimal("0"), "iva": Decimal("0")}
        if not is_nc:
            por_cliente[cid]["bruto"] += d.total
            por_cliente[cid]["liquido"] += d.total
            por_cliente[cid]["iva"] += d.total_iva
        else:
            por_cliente[cid]["liquido"] -= d.total
            por_cliente[cid]["iva"] -= d.total_iva

    return {
        "totais": {"bruto": bruto, "liquido": liquido, "iva": total_iva, "documentos": len(docs)},
        "por_cliente": [
            {"cliente_id": k, "nome": v["nome"], "bruto": v["bruto"], "liquido": v["liquido"], "iva": v["iva"]}
            for k, v in por_cliente.items() if v["bruto"] > 0 or v["liquido"] > 0
        ],
    }


def relatorio_documentos(
    db: Session,
    data_inicio: date,
    data_fim: date,
    tipo: str | None = None,
    serie_id: int | None = None,
    estado: str | None = None,
):
    q = db.query(Documento).filter(
        Documento.data_documento >= data_inicio,
        Documento.data_documento <= data_fim,
    )
    if tipo:
        q = q.filter(Documento.tipo_documento == tipo)
    if serie_id:
        q = q.filter(Documento.serie_id == serie_id)
    if estado:
        q = q.filter(Documento.estado == estado)
    return q.order_by(Documento.data_documento, Documento.id).all()


def relatorio_numeracao(db: Session):
    series = db.query(Serie).filter(Serie.ativo.is_(True)).all()
    resultados = []
    for s in series:
        if s.ultimo_numero == 0:
            continue
        numeros = db.query(Documento.numero).filter(
            Documento.serie_id == s.id,
            Documento.estado != EstadoDocumento.RASCUNHO,
        ).all()
        emitidos = {r[0] for r in numeros if r[0] > 0}
        esperados = set(range(1, s.ultimo_numero + 1))
        faltas = sorted(esperados - emitidos)
        resultados.append({
            "serie": s.prefixo,
            "tipo": s.tipo_documento.value,
            "ultimo": s.ultimo_numero,
            "emitidos": len(emitidos),
            "faltas": faltas,
        })
    return resultados


def relatorio_saldos(db: Session):
    docs = db.query(Documento).filter(
        Documento.estado == EstadoDocumento.FINALIZADO,
        Documento.tipo_documento.in_([
            TipoDocumentoFiscal.FT,
            TipoDocumentoFiscal.VD,
            TipoDocumentoFiscal.FS,
        ]),
    ).all()

    hoje = date.today()
    escaloes = {"0-30": Decimal("0"), "31-60": Decimal("0"), "61-90": Decimal("0"), "91-180": Decimal("0"), ">180": Decimal("0")}
    por_cliente: dict[int, dict] = {}

    for d in docs:
        pago = db.query(func.coalesce(func.sum(Pagamento.valor), 0)).filter(Pagamento.documento_id == d.id).scalar()
        creditado = _calcular_credito(db, d.id)
        saldo = Decimal(d.total) - Decimal(pago) - creditado
        if saldo <= 0:
            continue

        dias = (hoje - d.data_documento).days
        if dias <= 30:
            faixa = "0-30"
        elif dias <= 60:
            faixa = "31-60"
        elif dias <= 90:
            faixa = "61-90"
        elif dias <= 180:
            faixa = "91-180"
        else:
            faixa = ">180"

        escaloes[faixa] += saldo
        cid = d.cliente_id or 0
        if cid not in por_cliente:
            por_cliente[cid] = {"nome": d.cliente_nome or "Consumidor Final", "saldo": Decimal("0"), "faixa": faixa, "dias": dias}
        por_cliente[cid]["saldo"] += saldo
        if dias > por_cliente[cid]["dias"]:
            por_cliente[cid]["faixa"] = faixa
            por_cliente[cid]["dias"] = dias

    return {
        "totais": escaloes,
        "por_cliente": sorted(list(por_cliente.values()), key=lambda x: x["saldo"], reverse=True),
    }


def calcular_aging_cliente(db: Session, cliente_id: int) -> list[Decimal]:
    """Devolve [por_vencer, vencido_0_30, vencido_31_60, vencido_60+] em MZN."""
    hoje = date.today()
    docs = db.query(Documento).filter(
        Documento.cliente_id == cliente_id,
        Documento.estado == EstadoDocumento.FINALIZADO,
        Documento.tipo_documento.in_([
            TipoDocumentoFiscal.FT,
            TipoDocumentoFiscal.VD,
            TipoDocumentoFiscal.FS,
        ]),
    ).all()

    aging = [Decimal("0"), Decimal("0"), Decimal("0"), Decimal("0")]
    for d in docs:
        pago = db.query(func.coalesce(func.sum(Pagamento.valor), 0)).filter(
            Pagamento.documento_id == d.id
        ).scalar()
        creditado = _calcular_credito(db, d.id)
        saldo = Decimal(d.total) - Decimal(pago) - creditado
        if saldo <= 0:
            continue
        venc = d.data_vencimento or d.data_documento
        dias = (hoje - venc).days
        if dias <= 0:
            aging[0] += saldo
        elif dias <= 30:
            aging[1] += saldo
        elif dias <= 60:
            aging[2] += saldo
        else:
            aging[3] += saldo
    return aging


def relatorio_extrato_cliente(
    db: Session,
    cliente_id: int,
    data_inicio: date,
    data_fim: date,
):
    cliente = db.get(Cliente, cliente_id)
    if not cliente:
        raise ValueError("Cliente nao encontrado")

    def _soma_antes(tipo_lista, sinal=1):
        v = db.query(func.coalesce(func.sum(Documento.total), 0)).filter(
            Documento.cliente_id == cliente_id,
            Documento.estado == EstadoDocumento.FINALIZADO,
            Documento.data_documento < data_inicio,
            Documento.tipo_documento.in_(tipo_lista),
        ).scalar()
        return Decimal(v) * sinal

    # <-- AQUI: Saldo Inicial do cliente + movimentos anteriores ao período
    saldo_inicial = cliente.saldo_inicial + (
        _soma_antes([TipoDocumentoFiscal.FT, TipoDocumentoFiscal.VD, TipoDocumentoFiscal.FS], 1)
        - _soma_antes([TipoDocumentoFiscal.NC], 1)
        - _soma_antes([TipoDocumentoFiscal.RC], 1)
    )

    docs = db.query(Documento).filter(
        Documento.cliente_id == cliente_id,
        Documento.estado == EstadoDocumento.FINALIZADO,
        Documento.data_documento >= data_inicio,
        Documento.data_documento <= data_fim,
    ).all()

    movs: list[dict] = []
    for d in docs:
        if d.tipo_documento in (TipoDocumentoFiscal.FT, TipoDocumentoFiscal.VD, TipoDocumentoFiscal.FS):
            movs.append({
                "data": d.data_documento,
                "tipo": d.tipo_documento.value,
                "numero": d.numero_formatado,
                "descricao": f"{d.tipo_documento.value} emitida",
                "debito": d.total,
                "credito": Decimal("0"),
            })
        elif d.tipo_documento == TipoDocumentoFiscal.NC:
            movs.append({
                "data": d.data_documento,
                "tipo": "NC",
                "numero": d.numero_formatado,
                "descricao": "Nota de credito",
                "debito": Decimal("0"),
                "credito": d.total,
            })
        elif d.tipo_documento == TipoDocumentoFiscal.RC:
            movs.append({
                "data": d.data_documento,
                "tipo": "RC",
                "numero": d.numero_formatado,
                "descricao": "Recibo",
                "debito": Decimal("0"),
                "credito": d.total,
            })

    movs.sort(key=lambda m: (m["data"], m["numero"]))
    saldo = saldo_inicial
    movimentos = []
    for m in movs:
        saldo = saldo + m["debito"] - m["credito"]
        movimentos.append({
            "data": m["data"],
            "tipo": m["tipo"],
            "numero": m["numero"],
            "descricao": m["descricao"],
            "debito": m["debito"],
            "credito": m["credito"],
            "saldo": saldo,
        })

    return {
        "cliente": {
            "id": cliente.id,
            "numero": cliente.numero,
            "numero_formatado": cliente.numero_formatado,
            "nome": cliente.nome,
            "nuit": cliente.nuit,
            "morada": cliente.morada,
            "saldo_inicial": cliente.saldo_inicial,
        },
        "periodo": {"inicio": data_inicio, "fim": data_fim},
        "saldo_inicial": saldo_inicial,
        "saldo_final": saldo,
        "movimentos": movimentos,
    }


def relatorio_extrato_clientes_geral(
    db: Session,
    data_inicio: date,
    data_fim: date,
):
    clientes = db.query(Cliente).all()
    resultados = []

    for cliente in clientes:
        def _soma_antes(tipo_lista, sinal=1):
            v = db.query(func.coalesce(func.sum(Documento.total), 0)).filter(
                Documento.cliente_id == cliente.id,
                Documento.estado == EstadoDocumento.FINALIZADO,
                Documento.data_documento < data_inicio,
                Documento.tipo_documento.in_(tipo_lista),
            ).scalar()
            return Decimal(v) * sinal

        # <-- AQUI: Saldo Inicial do cliente + movimentos anteriores ao período
        saldo_inicial = cliente.saldo_inicial + (
            _soma_antes([TipoDocumentoFiscal.FT, TipoDocumentoFiscal.VD, TipoDocumentoFiscal.FS], 1)
            - _soma_antes([TipoDocumentoFiscal.NC], 1)
            - _soma_antes([TipoDocumentoFiscal.RC], 1)
        )

        docs_cliente = db.query(Documento).filter(
            Documento.cliente_id == cliente.id,
            Documento.estado == EstadoDocumento.FINALIZADO,
            Documento.data_documento >= data_inicio,
            Documento.data_documento <= data_fim,
        ).order_by(Documento.data_documento, Documento.id).all()

        movimentos = []
        total_debito = Decimal("0")
        total_credito = Decimal("0")

        for d in docs_cliente:
            if d.tipo_documento in (TipoDocumentoFiscal.FT, TipoDocumentoFiscal.VD, TipoDocumentoFiscal.FS):
                movimentos.append({
                    "tipo": d.tipo_documento.value,
                    "numero": d.numero_formatado,
                    "data": d.data_documento,
                    "debito": d.total,
                    "credito": Decimal("0"),
                })
                total_debito += d.total
            elif d.tipo_documento == TipoDocumentoFiscal.NC:
                movimentos.append({
                    "tipo": "NC",
                    "numero": d.numero_formatado,
                    "data": d.data_documento,
                    "debito": Decimal("0"),
                    "credito": d.total,
                })
                total_credito += d.total
            elif d.tipo_documento == TipoDocumentoFiscal.RC:
                movimentos.append({
                    "tipo": "RC",
                    "numero": d.numero_formatado,
                    "data": d.data_documento,
                    "debito": Decimal("0"),
                    "credito": d.total,
                })
                total_credito += d.total

        saldo_final = saldo_inicial + total_debito - total_credito

        if saldo_inicial != 0 or total_debito != 0 or total_credito != 0:
            resultados.append({
                "cliente_id": cliente.id,
                "numero_formatado": cliente.numero_formatado,
                "nome": cliente.nome,
                "nuit": cliente.nuit,
                "saldo_inicial": saldo_inicial,
                "total_debito": total_debito,
                "total_credito": total_credito,
                "saldo_final": saldo_final,
                "movimentos": movimentos,
            })

    total_debito_geral = sum(
        (c["total_debito"] for c in resultados), Decimal("0")
    )
    total_credito_geral = sum(
        (c["total_credito"] for c in resultados), Decimal("0")
    )
    saldo_final_geral = sum(
        (c["saldo_final"] for c in resultados), Decimal("0")
    )

    return {
        "periodo": {"inicio": data_inicio, "fim": data_fim},
        "clientes": sorted(resultados, key=lambda x: x["nome"]),
        "totais": {
            "debito": total_debito_geral,
            "credito": total_credito_geral,
            "saldo_final": saldo_final_geral,
        },
    }

    return {
        "periodo": {"inicio": data_inicio, "fim": data_fim},
        "clientes": sorted(resultados, key=lambda x: x["nome"]),
        "totais": {
            "debito": total_debito_geral,
            "credito": total_credito_geral,
            "saldo_final": saldo_final_geral,
        },
    }


# ==========================================
# EXPORTAÇÃO (CSV / Excel / PDF Genérico)
# ==========================================

def gerar_csv(dados: list[dict], headers: list[str]) -> bytes:
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=headers)
    writer.writeheader()
    writer.writerows(dados)
    return output.getvalue().encode("utf-8")


def gerar_excel(dados: list[dict], headers: list[str], titulo: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = titulo[:30]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = cell.font.copy(bold=True)
    for row in dados:
        ws.append([row.get(h, "") for h in headers])

    for idx, col_name in enumerate(headers, start=1):
        max_len = len(col_name)
        for row_idx in range(2, ws.max_row + 1):
            val = ws.cell(row=row_idx, column=idx).value
            if val:
                max_len = max(max_len, len(str(val)))
        col_letter = get_column_letter(idx)
        ws.column_dimensions[col_letter].width = max_len + 2

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def gerar_pdf_relatorio(dados: list[dict], headers: list[str], titulo: str) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4),
        leftMargin=10 * mm, rightMargin=10 * mm, topMargin=10 * mm, bottomMargin=10 * mm,
    )
    elements: list[Any] = []
    styles = getSampleStyleSheet()
    elements.append(Paragraph(titulo, styles["Heading1"]))

    table_data = [headers]
    for row in dados:
        table_data.append([str(row.get(h, "")) for h in headers])

    t = Table(table_data, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F5BA8")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("ALIGN", (0, 0), (-1, -1), "LEFT"),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#B8C4D4")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F6F8FB")]),
    ]))
    elements.append(t)
    doc.build(elements)
    return buf.getvalue()